import sys,json,zipfile,re
from pathlib import Path
from xml.etree import ElementTree as ET
sys.stdout.reconfigure(encoding='utf-8')
kind,path,limit=sys.argv[1],Path(sys.argv[2]),int(sys.argv[3])
parts=[];truncated=False;units=0

def add(text):
 global truncated
 current=sum(len(v) for v in parts)
 if current>=limit:truncated=True;return False
 text=str(text)
 if len(text)>limit-current:text=text[:limit-current];truncated=True
 parts.append(text);return not truncated

def archive():
 z=zipfile.ZipFile(path);infos=z.infolist()
 if len(infos)>3000 or sum(i.file_size for i in infos)>100*1024*1024:raise ValueError('文档解压后内容过大')
 if any(i.file_size>25*1024*1024 for i in infos):raise ValueError('文档内单项内容过大')
 return z

try:
 if kind=='pdf':
  from pypdf import PdfReader
  reader=PdfReader(str(path),strict=False)
  if reader.is_encrypted and not reader.decrypt(''):raise ValueError('PDF已加密，需要先提供可读取的文件')
  units=len(reader.pages)
  for i,page in enumerate(reader.pages[:100]):
   if not add('[第 %d 页]\n%s'%(i+1,page.extract_text() or '')):break
  if units>100:truncated=True
 elif kind in ('docx','pptx'):
  with archive() as z:
   if kind=='docx':
    names=['word/document.xml']
    if names[0] not in z.namelist():raise ValueError('不是有效的DOCX文档')
   else:
    names=sorted([n for n in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml',n)],key=lambda n:int(re.search(r'\d+',n).group()))
    if not names:raise ValueError('不是有效的PPTX文档')
   units=len(names)
   for name in names:
    root=ET.fromstring(z.read(name));text='\n'.join(e.text for e in root.iter() if e.tag.endswith('}t') and e.text)
    if not add(text):break
 elif kind=='xlsx':
  with archive() as z:
   if 'xl/workbook.xml' not in z.namelist():raise ValueError('不是有效的XLSX表格')
  import openpyxl
  book=openpyxl.load_workbook(path,read_only=True,data_only=True,keep_links=False);units=len(book.worksheets)
  for sheet in book.worksheets[:20]:
   if not add('[工作表: %s]'%sheet.title):break
   for row in sheet.iter_rows(max_row=min(sheet.max_row or 1000,1000),max_col=min(sheet.max_column or 50,50),values_only=True):
    if not add('\t'.join('' if v is None else str(v) for v in row)):break
   if (sheet.max_row or 0)>1000 or (sheet.max_column or 0)>50:truncated=True
  if units>20:truncated=True
  book.close()
 elif kind=='zip':
  with archive() as z:
   units=len(z.infolist())
   for info in z.infolist():
    if not add(json.dumps({'name':info.filename,'bytes':info.file_size},ensure_ascii=False)):break
 else:raise ValueError('不支持的文档解析类型')
 text='\n'.join(parts).strip()
 # Page markers alone do not count as a readable text layer.
 if kind=='pdf' and not re.sub(r'\[第 \d+ 页\]','',text).strip():text=''
 print(json.dumps({'text':text,'truncated':truncated,'units':units},ensure_ascii=False))
except Exception as error:
 print(json.dumps({'error':str(error)[:200]},ensure_ascii=False));sys.exit(2)
